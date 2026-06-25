"""Import SAP2000 hinge exports and compare them with proxy hinge estimates.

This module is intentionally independent from the SAP2000 COM layer. SAP2000
v22 installations that do not expose hinge result tables through OAPI can
still be calibrated by exporting a small validation subset from Show Tables.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Iterable


LEVELS = ("A-B", "B-IO", "IO-LS", "LS-CP", "CP-C", "C-D", "D-E", "beyond E")
LEVEL_RANK = {level: index for index, level in enumerate(LEVELS)}
VALIDATION_DIR = "hinge_validation"
RAW_EXPORT_DIR = "raw_exports"
NORMALIZED_DIR = "normalized"
REPORT_DIR = "reports"
PROXY_EVENTS_DIR = "proxy_events"
CALIBRATION_FILE = "proxy_calibration.json"

ALIASES = {
    "model": ("model", "model_name", "modelid", "model_id"),
    "case": ("case", "outputcase", "output_case", "loadcase", "load_case", "analysiscase", "analysis_case"),
    "step_number": ("stepnum", "step_num", "stepnumber", "step_number", "step", "stepno"),
    "load_step": ("steptype", "step_type", "loadstep", "load_step"),
    "element_name": ("frame", "framename", "frame_name", "frameobject", "frame_object", "object", "element", "element_name"),
    "hinge_name": ("hinge", "hingename", "hinge_name", "hingeproperty", "hinge_property", "assignhinge", "genhinge"),
    "hinge_location": ("location", "hingelocation", "hinge_location", "end", "stationtype", "station_type"),
    "relative_distance": ("reldist", "rel_dist", "relativedistance", "relative_distance", "station", "relative_location"),
    "hinge_state_level": ("state", "hingestate", "hinge_state", "status", "performancelevel", "performance_level"),
    "plastic_rotation_rad": ("rotation", "plasticrotation", "plastic_rotation", "plastic_rotation_rad", "deformation", "rot", "r3plastic", "r2plastic", "r1plastic"),
    "moment_kn_m": ("moment", "m", "m3", "moment_kn_m"),
    "axial_force_kn": ("axial", "p", "axialforce", "axial_force", "axial_force_kn"),
    "shear_v2_kn": ("v2", "shearv2", "shear_v2", "shear_v2_kn"),
    "shear_v3_kn": ("v3", "shearv3", "shear_v3", "shear_v3_kn"),
}


@dataclass(frozen=True)
class NormalizedHingeRow:
    """One normalized SAP2000 hinge result row."""

    model: str
    case: str
    direction: str
    step_number: float | None
    load_step: str
    element_name: str
    element_type: str
    hinge_name: str
    hinge_location: str
    relative_distance: float | None
    hinge_state_level: str
    plastic_rotation_rad: float | None
    moment_kn_m: float | None
    axial_force_kn: float | None
    shear_v2_kn: float | None
    shear_v3_kn: float | None


def select_validation_subset(output_dir: Path, count: int = 30) -> dict[str, Any]:
    """Select a deterministic, diverse model subset for manual SAP validation."""
    metadata_files = sorted(output_dir.glob("*_metadata.json"))
    candidates: list[dict[str, Any]] = []
    for path in metadata_files:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        if data.get("status") != "success":
            continue
        candidate = data.get("candidate", {})
        plastic = data.get("plastic_hinges", {}).get("results", {})
        summaries = plastic.get("summary_by_direction", {}) if isinstance(plastic, dict) else {}
        if not isinstance(candidate, dict) or not isinstance(summaries, dict):
            continue
        model_path = Path(str(data.get("model_path", "")))
        candidates.append(
            {
                "model": model_path.stem or path.stem.replace("_metadata", ""),
                "model_path": str(model_path),
                "metadata_path": str(path),
                "story_count": candidate.get("story_count"),
                "x_bay_count": candidate.get("x_bay_count"),
                "y_bay_count": candidate.get("y_bay_count"),
                "concrete_class": candidate.get("concrete_class"),
                "column_section": _section_label(candidate.get("column_section")),
                "beam_section": _section_label(candidate.get("beam_section")),
                "rho_col": candidate.get("rho_col"),
                "rho_beam_top": candidate.get("beam_top_ratio_support"),
                "rho_beam_bottom": candidate.get("beam_bottom_ratio_span"),
                "target_drift": candidate.get("pushover_target_drift_ratio"),
                "soil_class": candidate.get("soil_class"),
                "proxy_x_critical_type": _critical_type(summaries.get("X")),
                "proxy_y_critical_type": _critical_type(summaries.get("Y")),
                "proxy_x_first_type": _first_type(summaries.get("X")),
                "proxy_y_first_type": _first_type(summaries.get("Y")),
            }
        )
    if not candidates:
        raise ValueError(f"No successful metadata JSON files found in {output_dir}.")

    selected = _diverse_pick(candidates, min(count, len(candidates)))
    validation_root = output_dir / VALIDATION_DIR
    raw_dir = validation_root / RAW_EXPORT_DIR
    normalized_dir = validation_root / NORMALIZED_DIR
    report_dir = validation_root / REPORT_DIR
    for path in (raw_dir, normalized_dir, report_dir):
        path.mkdir(parents=True, exist_ok=True)

    manifest_path = validation_root / "validation_subset.csv"
    _write_csv(manifest_path, selected)
    instructions_path = validation_root / "README.txt"
    instructions_path.write_text(_instructions_text(raw_dir), encoding="utf-8")
    return {
        "selected_count": len(selected),
        "manifest_path": str(manifest_path),
        "raw_export_dir": str(raw_dir),
        "normalized_dir": str(normalized_dir),
        "report_dir": str(report_dir),
        "instructions_path": str(instructions_path),
    }


def import_hinge_exports(output_dir: Path, source_dir: Path | None = None) -> dict[str, Any]:
    """Normalize manually exported SAP2000 hinge CSV/XLSX files."""
    validation_root = output_dir / VALIDATION_DIR
    source_dir = source_dir or validation_root / RAW_EXPORT_DIR
    normalized_dir = validation_root / NORMALIZED_DIR
    normalized_dir.mkdir(parents=True, exist_ok=True)
    files = sorted(path for path in source_dir.iterdir() if path.suffix.lower() in {".csv", ".xlsx"})
    if not files:
        raise ValueError(f"No CSV/XLSX exports found in {source_dir}.")
    imported: list[dict[str, Any]] = []
    all_rows: list[dict[str, Any]] = []
    for path in files:
        rows = normalize_export_file(path)
        if not rows:
            imported.append({"file": path.name, "row_count": 0, "warning": "No recognizable hinge rows."})
            continue
        model = _infer_model_name(path, rows)
        clean_rows = [{**asdict(row), "model": row.model or model} for row in rows]
        target = normalized_dir / f"{model}_hinges_normalized.csv"
        _write_csv(target, clean_rows)
        all_rows.extend(clean_rows)
        imported.append({"file": path.name, "model": model, "row_count": len(clean_rows), "normalized_path": str(target)})
    combined_path = normalized_dir / "all_hinges_normalized.csv"
    _write_csv(combined_path, all_rows)
    return {"source_dir": str(source_dir), "imported": imported, "row_count": len(all_rows), "combined_path": str(combined_path)}


def compare_proxy_with_exact(output_dir: Path) -> dict[str, Any]:
    """Compare imported exact SAP hinge rows with stored proxy event rows."""
    normalized_path = output_dir / VALIDATION_DIR / NORMALIZED_DIR / "all_hinges_normalized.csv"
    if not normalized_path.exists():
        raise ValueError("Run the import command first; all_hinges_normalized.csv is missing.")
    exact_rows = list(csv.DictReader(normalized_path.open("r", encoding="utf-8-sig", newline="")))
    proxy_by_key = _load_proxy_rows(output_dir)
    comparison_rows: list[dict[str, Any]] = []
    state_pairs: Counter[tuple[str, str]] = Counter()
    matched = 0
    rotation_errors: list[float] = []
    for exact in exact_rows:
        key = _coarse_match_key(exact)
        proxy = _best_proxy_match(exact, proxy_by_key.get(key, []))
        row = dict(exact)
        row["proxy_matched"] = bool(proxy)
        if proxy:
            matched += 1
            exact_state = normalize_level(exact.get("hinge_state_level", ""))
            proxy_state = normalize_level(proxy.get("hinge_state_level", ""))
            row["proxy_hinge_state_level"] = proxy_state
            row["proxy_element_type"] = proxy.get("element_type")
            row["proxy_hinge_dof"] = proxy.get("hinge_dof")
            row["state_rank_delta"] = LEVEL_RANK.get(proxy_state, -1) - LEVEL_RANK.get(exact_state, -1)
            state_pairs[(exact_state, proxy_state)] += 1
            proxy_rotation = _float(proxy.get("plastic_rotation_rad"))
            exact_rotation = _float(exact.get("plastic_rotation_rad"))
            row["proxy_plastic_rotation_rad"] = proxy_rotation
            if proxy_rotation is not None and exact_rotation is not None:
                error = proxy_rotation - exact_rotation
                row["rotation_error_rad"] = error
                rotation_errors.append(error)
        comparison_rows.append(row)

    report_dir = output_dir / VALIDATION_DIR / REPORT_DIR
    report_dir.mkdir(parents=True, exist_ok=True)
    comparison_path = report_dir / "hinge_proxy_comparison.csv"
    _write_csv(comparison_path, comparison_rows)
    summary = {
        "exact_row_count": len(exact_rows),
        "matched_row_count": matched,
        "match_rate": round(matched / len(exact_rows), 4) if exact_rows else 0.0,
        "state_exact_match_count": sum(count for (exact, proxy), count in state_pairs.items() if exact == proxy),
        "state_exact_match_rate": round(
            sum(count for (exact, proxy), count in state_pairs.items() if exact == proxy) / matched, 4
        )
        if matched
        else 0.0,
        "rotation_mae_rad": round(sum(abs(value) for value in rotation_errors) / len(rotation_errors), 6) if rotation_errors else None,
        "rotation_bias_rad": round(sum(rotation_errors) / len(rotation_errors), 6) if rotation_errors else None,
        "state_confusion": _nested_counter(state_pairs),
        "comparison_csv": str(comparison_path),
    }
    summary_path = report_dir / "hinge_proxy_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown_path = report_dir / "hinge_proxy_summary.md"
    markdown_path.write_text(_summary_markdown(summary), encoding="utf-8")
    return {**summary, "summary_json": str(summary_path), "summary_markdown": str(markdown_path)}


def calibrate_proxy(output_dir: Path) -> dict[str, Any]:
    """Fit robust plastic-rotation scale factors from exact SAP hinge rows."""
    report_dir = output_dir / VALIDATION_DIR / REPORT_DIR
    comparison_path = report_dir / "hinge_proxy_comparison.csv"
    if not comparison_path.exists():
        raise ValueError("Run the report command first; hinge_proxy_comparison.csv is missing.")
    rows = list(csv.DictReader(comparison_path.open("r", encoding="utf-8-sig", newline="")))
    grouped: dict[str, list[float]] = defaultdict(list)
    all_ratios: list[float] = []
    state_boundaries: dict[str, list[float]] = defaultdict(list)
    state_boundary_samples: Counter[str] = Counter()
    for row in rows:
        if str(row.get("proxy_matched", "")).lower() not in {"true", "1"}:
            continue
        exact_rotation = _float(row.get("plastic_rotation_rad"))
        proxy_rotation = _float(row.get("proxy_plastic_rotation_rad"))
        exact_state = normalize_level(str(row.get("hinge_state_level", "")))
        if proxy_rotation is not None and proxy_rotation >= 0.0 and exact_state:
            for level in LEVELS[1:]:
                boundary_key = _boundary_key(row, level)
                if LEVEL_RANK.get(exact_state, 0) >= LEVEL_RANK[level]:
                    state_boundaries[boundary_key].append(proxy_rotation)
                    state_boundaries[_fallback_boundary_key(row, level)].append(proxy_rotation)
                    state_boundaries[f"all:all:{level}"].append(proxy_rotation)
                    state_boundary_samples[boundary_key] += 1
        if exact_rotation is None or proxy_rotation is None or proxy_rotation <= 1.0e-9:
            continue
        ratio = exact_rotation / proxy_rotation
        if not math.isfinite(ratio) or ratio <= 0.0:
            continue
        ratio = min(4.0, max(0.25, ratio))
        element_type = str(row.get("proxy_element_type") or row.get("element_type") or "all").lower()
        hinge_dof = str(row.get("proxy_hinge_dof") or _hinge_dof_from_name(str(row.get("hinge_name", ""))) or "all").upper()
        grouped[f"{element_type}:{hinge_dof}"].append(ratio)
        grouped[f"{element_type}:all"].append(ratio)
        all_ratios.append(ratio)
    if not all_ratios and not state_boundaries:
        raise ValueError("No matched rows with usable exact/proxy plastic rotations or states were found.")
    factors = {
        key: {"rotation_scale": round(_median(values), 6), "sample_count": len(values)}
        for key, values in sorted(grouped.items())
    }
    default_scale = round(_median(all_ratios), 6) if all_ratios else 1.0
    calibration = {
        "version": 1,
        "source_comparison_csv": str(comparison_path),
        "row_count": len(rows),
        "usable_rotation_count": len(all_ratios),
        "default": {"rotation_scale": default_scale, "sample_count": len(all_ratios)},
        "factors": factors,
        "state_boundaries": _state_boundary_payload(state_boundaries, state_boundary_samples),
        "note": "Plastic rotation scales are robust medians of clipped exact_SAP_rotation / proxy_rotation ratios. State boundaries are learned from the proxy rotation values where exact SAP hinge rows reached each state.",
    }
    calibration_path = output_dir / VALIDATION_DIR / CALIBRATION_FILE
    calibration_path.write_text(json.dumps(calibration, ensure_ascii=False, indent=2), encoding="utf-8")
    return {**calibration, "calibration_path": str(calibration_path)}


def normalize_export_file(path: Path) -> list[NormalizedHingeRow]:
    """Read one CSV/XLSX file and normalize recognizable hinge rows."""
    raw_rows = _read_table(path)
    model = _model_from_filename(path.name)
    file_direction = _direction_from_filename(path.name)
    rows: list[NormalizedHingeRow] = []
    for raw in raw_rows:
        index = {_normalize_header(key): value for key, value in raw.items() if key is not None}
        element = str(_pick(index, "element_name") or "").strip()
        state = normalize_level(str(_pick(index, "hinge_state_level") or ""))
        if not element or not state:
            continue
        case = str(_pick(index, "case") or "").strip()
        location = str(_pick(index, "hinge_location") or "").strip()
        relative = _float(_pick(index, "relative_distance"))
        hinge_name = str(_pick(index, "hinge_name") or "")
        rows.append(
            NormalizedHingeRow(
                model=str(_pick(index, "model") or model),
                case=case,
                direction=_direction(case) or file_direction or "X",
                step_number=_float(_pick(index, "step_number")),
                load_step=str(_pick(index, "load_step") or ""),
                element_name=element,
                element_type=_element_type(element),
                hinge_name=hinge_name,
                hinge_location=location or _location_from_relative(relative),
                relative_distance=relative,
                hinge_state_level=state,
                plastic_rotation_rad=_hinge_rotation(index, hinge_name),
                moment_kn_m=_hinge_moment(index, hinge_name),
                axial_force_kn=_float(_pick(index, "axial_force_kn")),
                shear_v2_kn=_float(_pick(index, "shear_v2_kn")),
                shear_v3_kn=_float(_pick(index, "shear_v3_kn")),
            )
        )
    return rows


def normalize_level(value: str) -> str:
    """Normalize common SAP hinge state spellings to the project state levels."""
    text = re.sub(r"\s+", "", str(value or "")).upper().replace("_", "-").replace("TO", "-")
    aliases = {
        "A-B": "A-B", "AB": "A-B", "B-IO": "B-IO", "BIO": "B-IO",
        "IO-LS": "IO-LS", "IOLS": "IO-LS", "LS-CP": "LS-CP", "LSCP": "LS-CP",
        "CP-C": "CP-C", "CPC": "CP-C", "C-D": "C-D", "CD": "C-D",
        "D-E": "D-E", "DE": "D-E", ">E": "beyond E", "BEYONDE": "beyond E",
    }
    return aliases.get(text, "")


def _diverse_pick(candidates: list[dict[str, Any]], count: int) -> list[dict[str, Any]]:
    """Greedily pick candidates that cover stories, drift, soil, and proxy outcomes."""
    selected: list[dict[str, Any]] = []
    remaining = sorted(candidates, key=lambda row: row["model"])
    covered: Counter[tuple[str, Any]] = Counter()
    dimensions = (
        "story_count", "target_drift", "soil_class", "rho_col",
        "proxy_x_critical_type", "proxy_y_critical_type", "proxy_x_first_type", "proxy_y_first_type",
    )
    while remaining and len(selected) < count:
        def score(row: dict[str, Any]) -> tuple[float, str]:
            novelty = sum(1.0 / (1.0 + covered[(field, row.get(field))]) for field in dimensions)
            return novelty, row["model"]
        best = max(remaining, key=score)
        remaining.remove(best)
        selected.append(best)
        for field in dimensions:
            covered[(field, best.get(field))] += 1
    return selected


def _load_proxy_rows(output_dir: Path) -> dict[tuple[str, str, str], list[dict[str, Any]]]:
    """Load compact proxy events from per-model JSON metadata files."""
    rows: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    proxy_dir = output_dir / VALIDATION_DIR / PROXY_EVENTS_DIR
    if proxy_dir.exists():
        for path in sorted(proxy_dir.glob("*_proxy_hinges.csv")):
            try:
                with path.open("r", encoding="utf-8-sig", newline="") as file:
                    for event in csv.DictReader(file):
                        event_model = Path(str(event.get("model") or _model_from_filename(path.name))).stem
                        event_direction = str(event.get("direction") or _direction(str(event.get("case", ""))) or _direction_from_filename(path.name) or "X").upper()
                        enriched = {**event, "model": event_model, "direction": event_direction}
                        rows[_coarse_match_key(enriched)].append(enriched)
            except OSError:
                continue
    for path in output_dir.glob("*_metadata.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        model = Path(str(data.get("model_path", ""))).stem or path.stem.replace("_metadata", "")
        events = data.get("plastic_hinges", {}).get("results", {}).get("proxy_events_by_direction", {})
        if not isinstance(events, dict):
            continue
        for direction, payload in events.items():
            if not isinstance(payload, dict):
                continue
            for event in payload.get("events", []):
                if isinstance(event, dict):
                    rows[_coarse_match_key({**event, "model": model, "direction": direction})].append(event)
    return rows


def _best_proxy_match(exact: dict[str, Any], options: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Return the proxy event nearest to an exact SAP result row."""
    if not options:
        return None
    exact_step = _float(exact.get("step_number"))
    exact_relative = _float(exact.get("relative_distance"))
    return min(
        options,
        key=lambda row: (
            abs((_float(row.get("step_number")) or 0.0) - (exact_step or 0.0)),
            abs((_float(row.get("relative_distance")) or 0.0) - (exact_relative or 0.0)),
        ),
    )


def _match_key(row: dict[str, Any]) -> tuple[str, str, str, str]:
    """Return model, direction, element, and hinge-location matching key."""
    model = Path(str(row.get("model", ""))).stem
    direction = str(row.get("direction") or _direction(str(row.get("case", "")))).upper()
    element = str(row.get("element_name", "")).strip()
    location = str(row.get("hinge_location") or _location_from_relative(_float(row.get("relative_distance")))).lower()
    return model, direction, element, location


def _coarse_match_key(row: dict[str, Any]) -> tuple[str, str, str]:
    """Return tolerant model, direction, and element matching key."""
    model, direction, element, _ = _match_key(row)
    return model, direction, element


def _read_table(path: Path) -> list[dict[str, Any]]:
    """Read CSV or XLSX records without imposing one SAP2000 locale."""
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1254", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", errors="replace")
        lines = text.splitlines()
        if lines and lines[0].lstrip("\ufeff").startswith("TABLE:"):
            lines = lines[1:]
        sample = "\n".join(lines[:10])
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        records = list(csv.DictReader(lines, delimiter=delimiter))
        return [row for row in records if not _is_sap_units_row(row)]
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX import requires openpyxl. Export CSV or install openpyxl.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        records = list(sheet.iter_rows(values_only=True))
        if not records:
            return []
        headers = [str(value or "") for value in records[0]]
        return [dict(zip(headers, values)) for values in records[1:]]
    return []


def _pick(index: dict[str, Any], name: str) -> Any:
    """Pick a normalized table value using known aliases."""
    for alias in ALIASES[name]:
        if _normalize_header(alias) in index:
            return index[_normalize_header(alias)]
    return None


def _is_sap_units_row(row: dict[str, Any]) -> bool:
    """Return true for SAP2000's column-type/unit row below table headers."""
    values = {str(value or "").strip().lower() for value in row.values()}
    return "text" in values and ("unitless" in values or "radians" in values or "kn" in values)


def _hinge_rotation(index: dict[str, Any], hinge_name: str) -> float | None:
    """Return plastic rotation matching the generated hinge degree of freedom."""
    dof = _hinge_dof_from_name(hinge_name)
    preferred = {"M2": "r2plastic", "M3": "r3plastic", "V2": "r2plastic", "V3": "r3plastic"}.get(dof)
    if preferred:
        value = _float(index.get(_normalize_header(preferred)))
        if value is not None:
            return value
    return _float(_pick(index, "plastic_rotation_rad"))


def _hinge_moment(index: dict[str, Any], hinge_name: str) -> float | None:
    """Return the SAP moment component associated with one hinge."""
    dof = _hinge_dof_from_name(hinge_name)
    preferred = {"M2": "m2", "M3": "m3"}.get(dof)
    if preferred:
        value = _float(index.get(_normalize_header(preferred)))
        if value is not None:
            return value
    return _float(_pick(index, "moment_kn_m"))


def _normalize_header(value: str) -> str:
    """Normalize CSV/XLSX header labels."""
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _float(value: Any) -> float | None:
    """Parse a locale-tolerant floating point value."""
    if value in (None, ""):
        return None
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def _direction(case: str) -> str:
    """Infer pushover direction from case name."""
    text = str(case or "").upper()
    if re.search(r"(^|[_ -])Y($|[_ -])", text) or text.endswith("Y"):
        return "Y"
    if re.search(r"(^|[_ -])X($|[_ -])", text) or text.endswith("X"):
        return "X"
    return ""


def _direction_from_filename(filename: str) -> str:
    """Infer pushover direction from an exported SAP table filename."""
    text = Path(filename).stem.lower()
    if re.search(r"(__|_)pushover[_-]?y($|__|_)", text):
        return "Y"
    if re.search(r"(__|_)pushover[_-]?x($|__|_)", text):
        return "X"
    return ""


def _element_type(name: str) -> str:
    """Infer generated element type from frame name."""
    return "column" if str(name).upper().startswith("C_") else "beam"


def _hinge_dof_from_name(name: str) -> str:
    """Infer hinge DOF from generated hinge property name."""
    match = re.search(r"(M2|M3|V2|V3)", str(name or "").upper())
    return match.group(1) if match else ""


def _median(values: list[float]) -> float:
    """Return median for a non-empty numeric list."""
    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) / 2.0


def _location_from_relative(value: float | None) -> str:
    """Convert relative hinge station to generated end label."""
    if value is None:
        return ""
    if math.isclose(value, 0.0, abs_tol=0.02):
        return "i-end"
    if math.isclose(value, 1.0, abs_tol=0.02):
        return "j-end"
    return f"relative-{value:.3f}"


def _model_from_filename(filename: str) -> str:
    """Infer generated model stem from export file name."""
    stem = Path(filename).stem
    stem = re.sub(r"__(frame_)?hinge_states$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"__pushover_[xy]$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"(_hinges?|_hinge_results?|_results?|_export)$", "", stem, flags=re.IGNORECASE)
    return stem


def _boundary_key(row: dict[str, Any], level: str) -> str:
    """Return a specific state-boundary calibration key."""
    element_type = str(row.get("proxy_element_type") or row.get("element_type") or "all").lower()
    hinge_dof = str(row.get("proxy_hinge_dof") or _hinge_dof_from_name(str(row.get("hinge_name", ""))) or "all").upper()
    return f"{element_type}:{hinge_dof}:{level}"


def _fallback_boundary_key(row: dict[str, Any], level: str) -> str:
    """Return an element-only state-boundary calibration key."""
    element_type = str(row.get("proxy_element_type") or row.get("element_type") or "all").lower()
    return f"{element_type}:all:{level}"


def _state_boundary_payload(boundaries: dict[str, list[float]], samples: Counter[str]) -> dict[str, dict[str, float | int]]:
    """Summarize learned proxy-rotation thresholds for hinge states."""
    payload: dict[str, dict[str, float | int]] = {}
    for key, values in sorted(boundaries.items()):
        clean = [value for value in values if value is not None and math.isfinite(value) and value >= 0.0]
        if not clean:
            continue
        payload[key] = {
            "rotation_threshold_rad": round(_percentile(clean, 0.10), 6),
            "median_rotation_rad": round(_median(clean), 6),
            "sample_count": len(clean),
            "specific_sample_count": int(samples.get(key, 0)),
        }
    return payload


def _percentile(values: list[float], fraction: float) -> float:
    """Return a simple nearest-rank percentile for non-empty values."""
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, round((len(ordered) - 1) * fraction)))
    return ordered[index]


def _infer_model_name(path: Path, rows: list[NormalizedHingeRow]) -> str:
    """Prefer an explicit model field and fall back to export filename."""
    for row in rows:
        if row.model:
            return Path(row.model).stem
    return _model_from_filename(path.name)


def _section_label(value: Any) -> str:
    """Format JSON section dimensions."""
    if not isinstance(value, dict):
        return ""
    return f"{round(float(value.get('width', 0)) * 100)}x{round(float(value.get('depth', 0)) * 100)}"


def _critical_type(summary: Any) -> str:
    """Return proxy critical element type."""
    if not isinstance(summary, dict):
        return ""
    events = summary.get("critical_events", [])
    return str(events[0].get("element_type", "")) if isinstance(events, list) and events else ""


def _first_type(summary: Any) -> str:
    """Return proxy first-hinge element type."""
    if not isinstance(summary, dict):
        return ""
    return str((summary.get("first_plastic_hinge") or {}).get("element_type", ""))


def _nested_counter(counter: Counter[tuple[str, str]]) -> dict[str, dict[str, int]]:
    """Return confusion counter in JSON-friendly shape."""
    result: dict[str, dict[str, int]] = {}
    for (exact, proxy), count in sorted(counter.items()):
        result.setdefault(exact, {})[proxy] = count
    return result


def _write_csv(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    """Write dictionaries to UTF-8 CSV."""
    records = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not records:
        path.write_text("", encoding="utf-8-sig")
        return
    fields: list[str] = []
    for row in records:
        for key in row:
            if key not in fields:
                fields.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


def _instructions_text(raw_dir: Path) -> str:
    """Return concise manual SAP export instructions."""
    return f"""SAP2000 v22 hinge validation export

1. Open one model listed in validation_subset.csv.
2. Run analysis if required.
3. Open Display > Show Tables.
4. Select the nonlinear static hinge-result table for PUSHOVER_X and PUSHOVER_Y.
5. Export CSV or XLSX into:
   {raw_dir}
6. Use the model stem as filename, for example:
   Model_0001_..._hinges.csv

Required columns:
- analysis case
- step number
- frame/element name
- hinge state

Recommended columns:
- hinge name
- hinge location or relative station
- plastic rotation/deformation
- moment, axial force, V2, V3

Then run:
python sap_hinge_validation.py import --output-dir "D:\\sap2000_generated_models"
python sap_hinge_validation.py report --output-dir "D:\\sap2000_generated_models"
python sap_hinge_validation.py calibrate --output-dir "D:\\sap2000_generated_models"
"""


def _summary_markdown(summary: dict[str, Any]) -> str:
    """Return a compact Markdown proxy validation report."""
    lines = [
        "# SAP2000 Hinge Proxy Validation",
        "",
        f"- Exact SAP rows: {summary['exact_row_count']}",
        f"- Matched proxy rows: {summary['matched_row_count']}",
        f"- Match rate: {summary['match_rate']:.2%}",
        f"- Exact hinge-state agreement: {summary['state_exact_match_rate']:.2%}",
        f"- Plastic rotation MAE (rad): {summary['rotation_mae_rad']}",
        f"- Plastic rotation bias (rad): {summary['rotation_bias_rad']}",
        "",
        "## State Confusion",
        "",
        "```json",
        json.dumps(summary["state_confusion"], ensure_ascii=False, indent=2),
        "```",
    ]
    return "\n".join(lines) + "\n"


def build_parser() -> argparse.ArgumentParser:
    """Build command-line interface."""
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    for command in ("select", "import", "report", "calibrate"):
        subparser = subparsers.add_parser(command)
        subparser.add_argument("--output-dir", type=Path, default=Path(r"D:\sap2000_generated_models"))
        if command == "select":
            subparser.add_argument("--count", type=int, default=30)
        if command == "import":
            subparser.add_argument("--source-dir", type=Path)
    return parser


def main() -> None:
    """Run one hinge-validation command."""
    args = build_parser().parse_args()
    if args.command == "select":
        result = select_validation_subset(args.output_dir, args.count)
    elif args.command == "import":
        result = import_hinge_exports(args.output_dir, args.source_dir)
    elif args.command == "calibrate":
        result = calibrate_proxy(args.output_dir)
    else:
        result = compare_proxy_with_exact(args.output_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
